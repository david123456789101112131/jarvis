"""Presentations: topic outlines, a self-contained HTML slide deck and a native PPTX file.

`outline_from_topic` writes a complete talk (title, agenda, background, key points, examples,
plan, summary, Q&A) in Hebrew or English; `build_presentation_html` turns slides into a
single-file deck with keyboard navigation, speaker notes and print-to-PDF styles;
`build_presentation_pptx` writes the same slides as an editable 16:9 PowerPoint file with
right-to-left paragraphs for Hebrew.
"""

from __future__ import annotations

import io
import json
import re
from collections.abc import Sequence
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Literal

from ._util import FONT_STACK, esc, has_hebrew, make_rng, resolve_lang, text_dir
from .assets import Asset
from .domains import GENERIC, Domain, detect_domain
from .palette import Palette, generate_palette, hex_to_rgb, mix

Layout = Literal["title", "section", "bullets", "two_column", "quote", "closing"]
LAYOUTS: tuple[str, ...] = ("title", "section", "bullets", "two_column", "quote", "closing")


@dataclass
class Slide:
    """One slide. Layout-specific use of `bullets`:

    title/closing – subtitle lines; section – an optional subtitle; quote – [quote, attribution];
    two_column – the first half fills the first column (headed by `columns[0]`), the rest the second.
    """

    title: str
    bullets: list[str] = field(default_factory=list)
    notes: str = ""
    layout: Layout = "bullets"
    columns: tuple[str, str] | None = None
    numbered: bool = False

    def __post_init__(self) -> None:
        if self.layout not in LAYOUTS:
            raise ValueError(f"unknown slide layout {self.layout!r}; choose one of {', '.join(LAYOUTS)}")
        self.bullets = [str(b) for b in (self.bullets or [])]

    def to_dict(self) -> dict:
        data = asdict(self)
        data["columns"] = list(self.columns) if self.columns else None
        return data

    @classmethod
    def from_dict(cls, data: dict) -> Slide:
        cols = data.get("columns")
        return cls(title=str(data.get("title", "")), bullets=list(data.get("bullets") or []),
                   notes=str(data.get("notes", "")), layout=data.get("layout", "bullets"),
                   columns=(str(cols[0]), str(cols[1])) if cols and len(cols) == 2 else None,
                   numbered=bool(data.get("numbered", False)))


# --------------------------------------------------------------------------------------------
# Outline
# --------------------------------------------------------------------------------------------

_EN_DECK = re.compile(
    r"^(?:please\s+)?(?:(?:make|create|build|prepare|write|generate|give|design)\s+(?:me\s+|us\s+)?)?(?:a\s+|an\s+|the\s+)?"
    r"(?:short\s+|quick\s+|\d+[- ]slides?\s+)?(?:presentation|deck|slide\s*deck|slides|talk|pitch(?:\s+deck)?|lecture)"
    r"\s+(?:about|on|for|regarding|covering|to)\s+(?:(?:a|an|the)\s+)?",
    re.I,
)
_HE_DECK = re.compile(
    r"^(?:ג'?ארוויס[,:]?\s+)?(?:בבקשה\s+)?(?:(?:תכין|הכן|תכיני|הכיני|צור|תצור|צרי|תכתוב|כתוב|תבנה|בנה|תעשה|"
    r"אני צריך|אני צריכה|אני רוצה)\s+(?:לי\s+|לנו\s+)?)?(?:מצגת|הרצאה|פרזנטציה|מצגות)\s+"
    r"(?:(?:קצרה|מקצועית)\s+)?(?:על\s+|בנושא\s+|עבור\s+|בשביל\s+|ל(?=[א-ת]))"
)
_BUSINESS_WORDS = (
    "עסק", "חנות", "מותג", "שיווק", "השקה", "השקת", "פתיחת", "תוכנית עסקית", "סטארטאפ", "חברה", "מוצר", "לקוחות", "מכירות", "משקיעים", "אתר", "אפליקציה",
    "business", "store", "shop", "brand", "marketing", "launch", "startup", "company", "product", "customers", "sales",
    "investors", "pitch", "website", "app",
)

_O = {
    "he": {
        "subtitle": "סקירה, תובנות ותוכנית פעולה", "for_audience": "מצגת עבור {audience}",
        "agenda": "על מה נדבר היום", "background": "רקע: התמונה הגדולה", "challenge": "האתגר",
        "key_points": "נקודות מפתח", "compare": "המצב היום מול המטרה", "compare_cols": ("היום", "המטרה"),
        "examples": "דוגמאות מהשטח", "quote": "המסר המרכזי", "metrics": "איך נמדוד הצלחה", "plan": "תוכנית פעולה",
        "summary": "סיכום", "qa": "שאלות ותשובות", "qa_lines": ["תודה רבה על ההקשבה!", "נשמח לשמוע שאלות, רעיונות והערות"],
        "section_1": "חלק א׳: להבין את התמונה", "section_2": "חלק ב׳: מהרעיון לביצוע",
        "default_aud": "הקהל שלנו", "everyone": "צוותים, לקוחות ושותפים",
        "background_generic": [
            "למה עכשיו – השינויים של השנים האחרונות יוצרים הזדמנויות חדשות",
            "מי מושפע – {aud_list}",
            "מה על הכף – זמן, כסף ואמון",
            "המטרה שלנו – להבין את התמונה, לבחור כיוון ולצאת לדרך",
        ],
        "background_business": [
            "ההזדמנות – הביקוש ל{essence} רק הולך וגדל",
            "הקהל – {aud_list}",
            "ההצעה שלנו – {offer}",
            "המטרה – להפוך לשם שאנשים אוהבים וממליצים עליו לחברים",
        ],
        "challenge_items": [
            "משאבים מוגבלים – זמן, תקציב ואנשים",
            "יותר מדי כיוונים – קשה להחליט מה קודם",
            "הציפיות עולות – והתחרות רק גוברת",
            "מדידה – לא תמיד ברור מה באמת עובד",
        ],
        "key_pool": [
            "מתחילים מהצורך – מה באמת חשוב {to_aud}?", "פשטות – פחות אפשרויות, יותר בהירות",
            "צעדים קטנים – מתקדמים בהדרגה ומודדים כל שלב", "שיתוף – מערבים את כל השותפים כבר מההתחלה",
            "נתונים – החלטות שמבוססות על עובדות, לא על תחושות", "עקביות – הרגלים קטנים שמצטברים לתוצאה גדולה",
            "גמישות – בודקים, לומדים ומשנים כיוון כשצריך", "חוויה – כל נקודת מגע צריכה להיות פשוטה וברורה",
        ],
        "compare_today": ["תהליכים ידניים ומפוזרים", "החלטות על סמך תחושת בטן", "תוצאות לא עקביות"],
        "compare_goal": ["תהליך אחד, ברור ופשוט", "החלטות שמבוססות על נתונים", "שיפור מדיד מחודש לחודש"],
        "example_pool": [
            "פיילוט של שבועיים – בדיקה של רעיון אחד עם מדד הצלחה ברור",
            "צוות קטן שהתמקד בשלושה יעדים בלבד – והשיג את כולם",
            "שינוי קטן בתהליך – שחסך שעות עבודה בכל שבוע",
            "סקר קצר בקרב {aud} – שחשף צורך שלא ראינו קודם",
            "השוואה בין שתי גרסאות (A/B) – שהכריעה ויכוח ארוך בתוך יום",
            "שיתוף פעולה עם שותף חיצוני – שקיצר את הדרך לתוצאה",
        ],
        "quotes": ["צעד קטן היום שווה יותר מתוכנית מושלמת מחר.", "מה שנמדד – משתפר.",
                   "פשוט זה לא קל. פשוט זו עבודה קשה שנעשתה היטב.", "הרעיון הוא רק ההתחלה. הביצוע הוא הכול."],
        "metrics_items": [
            "מדד מוביל אחד – מספר אחד שכולם מכירים ובודקים", "בדיקה שבועית – רבע שעה: מה עבד ומה לא",
            "יעד ל-90 יום – עם נקודות ביקורת בכל חודש", "משוב – שיחות קבועות עם {aud}",
        ],
        "metrics_business_tail": "יעד ל-90 יום – עם נקודות ביקורת בכל חודש",
        "kpi_tracking": ["מעקב יומי בלוח הבקרה", "השוואה חודשית מול היעד", "בדיקת מגמה בכל רבעון"],
        "plan_items": [
            "שבועות 1–2 – מיפוי המצב הקיים והגדרת יעדים", "שבועות 3–4 – פיילוט קטן ומדיד",
            "חודש 2 – הרחבה על סמך התוצאות", "חודש 3 – הטמעה מלאה וסיכום תובנות",
        ],
        "summary_items": [
            "{topic} – הזדמנות, לא רק אתגר", "מיקוד בכמה נקודות מפתח עושה את ההבדל",
            "מתחילים בקטן, מודדים ומשפרים", "הצעד הבא – {next_step}",
        ],
        "notes": {
            "title": "פתיחה: הציגו את עצמכם ואת מטרת המצגת בשני משפטים, וספרו לקהל מה ייקחו איתם בסוף.",
            "agenda": "עברו בקצרה על סדר היום, כדי שהקהל יידע לאן אנחנו הולכים.",
            "background": "תנו הקשר: למה הנושא של {topic} חשוב דווקא עכשיו? אפשר לפתוח בנתון מפתיע או בסיפור קצר.",
            "challenge": "תארו את הקושי בצורה מוחשית, ושאלו את הקהל אם זה מוכר להם.",
            "key_points": "התעכבו כחצי דקה על כל נקודה, עם דוגמה קצרה לכל אחת.",
            "compare": "הדגישו את הפער בין המצב היום למטרה – זה הלב של המצגת.",
            "examples": "ספרו סיפור אחד בהרחבה; את השאר אפשר להזכיר בקצרה.",
            "quote": "עצרו לרגע של שקט אחרי המשפט. זה המסר שהקהל צריך לזכור.",
            "metrics": "הסבירו איך נדע שהצלחנו, מתי נבדוק ומי אחראי על המדידה.",
            "plan": "הציגו את לוח הזמנים ואת האחראים לכל שלב, ובקשו מחויבות לצעד הראשון.",
            "summary": "חזרו על שלוש הנקודות החשובות ביותר – במילים פשוטות.",
            "section": "מעבר: סכמו במשפט את מה שראינו עד עכשיו, והציגו את החלק הבא.",
            "qa": "הזמינו שאלות. אם אין, פתחו בשאלה שנשאלתם בעבר על הנושא של {topic}.",
        },
    },
    "en": {
        "subtitle": "Overview, insights and an action plan", "for_audience": "A presentation for {audience}",
        "agenda": "Agenda", "background": "Background: the big picture", "challenge": "The challenge",
        "key_points": "Key points", "compare": "Where we are vs. where we're going", "compare_cols": ("Today", "The goal"),
        "examples": "Real-world examples", "quote": "The key message", "metrics": "How we'll measure success",
        "plan": "Action plan", "summary": "Summary", "qa": "Questions & answers",
        "qa_lines": ["Thank you for listening!", "Questions, ideas and feedback are welcome"],
        "section_1": "Part 1: Understanding the picture", "section_2": "Part 2: From idea to execution",
        "default_aud": "our audience", "everyone": "teams, customers and partners",
        "background_generic": [
            "Why now: recent changes are opening up new opportunities",
            "Who it affects: {aud_list}",
            "What's at stake: time, money and trust",
            "Our goal: understand the picture, choose a direction and get moving",
        ],
        "background_business": [
            "The opportunity: demand for {essence} keeps growing",
            "The audience: {aud_list}",
            "Our offer: {offer}",
            "The goal: a {label} people love and recommend to friends",
        ],
        "challenge_items": [
            "Limited resources: time, budget and people",
            "Too many directions: it's hard to decide what comes first",
            "Rising expectations: and the competition keeps growing",
            "Measurement: it isn't always clear what really works",
        ],
        "key_pool": [
            "Start with the need: what really matters to {aud}?", "Keep it simple: fewer options, more clarity",
            "Small steps: move forward gradually and measure each stage",
            "Involve people: bring every stakeholder in from day one",
            "Use data: decisions based on facts, not hunches", "Be consistent: small habits add up to big results",
            "Stay flexible: test, learn and change course when needed",
            "Mind the experience: every touchpoint should be simple and clear",
        ],
        "compare_today": ["Manual, scattered processes", "Decisions based on gut feeling", "Inconsistent results"],
        "compare_goal": ["One clear, simple process", "Data-informed decisions", "Measurable improvement every month"],
        "example_pool": [
            "A two-week pilot: one idea tested against one clear success metric",
            "A small team that focused on just three goals, and hit all of them",
            "One small process change that saved hours every week",
            "A short survey of {aud} that revealed a need nobody had noticed",
            "An A/B comparison that settled a long debate within a day",
            "A partnership with an outside expert that shortened the path to results",
        ],
        "quotes": ["A small step today beats a perfect plan tomorrow.", "What gets measured gets improved.",
                   "Simple isn't easy. Simple is hard work done well.",
                   "An idea is only the beginning. Execution is everything."],
        "metrics_items": [
            "One leading metric: a single number everyone knows and checks",
            "A weekly check-in: fifteen minutes on what worked and what didn't",
            "A 90-day target: with checkpoints every month", "Feedback: regular conversations with {aud}",
        ],
        "metrics_business_tail": "A 90-day target: with checkpoints every month",
        "kpi_tracking": ["tracked daily on the dashboard", "compared monthly against the target",
                         "reviewed as a quarterly trend"],
        "plan_items": [
            "Weeks 1–2: map the current state and set goals", "Weeks 3–4: run a small, measurable pilot",
            "Month 2: expand based on the results", "Month 3: full rollout and a lessons-learned review",
        ],
        "summary_items": [
            "{Topic}: an opportunity, not just a challenge", "Focusing on a few key points makes the difference",
            "Start small, measure and improve", "Next step: {next_step}",
        ],
        "notes": {
            "title": "Opening: introduce yourself and the goal of the talk in two sentences, and tell the audience "
                     "what they'll walk away with.",
            "agenda": "Briefly walk through the agenda so everyone knows where we're heading.",
            "background": "Set the context for {topic}: why does it matter right now? Open with a surprising number "
                          "or a short story.",
            "challenge": "Describe the pain concretely, and ask the audience whether it sounds familiar.",
            "key_points": "Spend about thirty seconds on each point, with a short example for each.",
            "compare": "Stress the gap between today and the goal: this is the heart of the talk.",
            "examples": "Tell one story in depth; the others can be mentioned briefly.",
            "quote": "Pause for a moment after this line. It's the message people should remember.",
            "metrics": "Explain how we'll know we succeeded, when we'll check, and who owns the numbers.",
            "plan": "Present the timeline and the owner of each stage, and ask for commitment to the first step.",
            "summary": "Repeat the three most important points in plain words.",
            "section": "Transition: sum up what we've covered in one sentence, then introduce the next part.",
            "qa": "Invite questions. If there are none, start with a question you're often asked on this topic.",
        },
    },
}

# Priority of each optional slide when the requested length is shorter than the full outline.
_PRIORITY = {"title": 0, "qa": 1, "key_points": 2, "agenda": 3, "background": 4, "plan": 5, "summary": 6,
             "examples": 7, "challenge": 8, "compare": 9, "quote": 10, "metrics": 11, "section_1": 12, "section_2": 13}
_ORDER = ("title", "agenda", "section_1", "background", "challenge", "key_points", "compare", "examples", "quote",
          "section_2", "metrics", "plan", "summary", "qa")


def clean_presentation_topic(topic: str) -> str:
    """'make me a presentation about remote work' -> 'remote work'; 'תכין לי מצגת על בינה מלאכותית' -> 'בינה מלאכותית'."""
    t = (topic or "").strip().strip(".!?")
    out = _HE_DECK.sub("", t, count=1) if has_hebrew(t) else _EN_DECK.sub("", t, count=1)
    return out.strip(" ,.:;-") or t


def _he_prefix(prefix: str, word: str) -> str:
    """Attach a one-letter Hebrew preposition, dropping the definite article ('ל' + 'הצוות' -> 'לצוות')."""
    if word.startswith("ה") and prefix in ("ל", "ב", "כ"):
        return prefix + word[1:]
    return prefix + word


def _join_list(items: Sequence[str], lang: str) -> str:
    items = [i for i in items if i]
    if len(items) <= 1:
        return items[0] if items else ""
    if lang == "he":
        if any(" ו" in i for i in items):  # items that already contain 'ו' read better as a plain list
            return ", ".join(items)
        return ", ".join(items[:-1]) + " ו" + items[-1]
    return ", ".join(items[:-1]) + ("," if len(items) > 2 else "") + " and " + items[-1]


def _is_business(topic: str, domain: Domain) -> bool:
    if domain is GENERIC:
        return False
    from ._util import keyword_in, norm

    hay = norm(topic)
    labels = (domain.label["he"], domain.label["en"])
    return any(keyword_in(hay, w, normalized=True) for w in (*_BUSINESS_WORDS, *labels))


def _head_text(item: str) -> tuple[str, str]:
    """Split 'Head – text' / 'Head: text' into its parts (head is empty when there is none)."""
    for sep in (" – ", ": "):
        head, found, rest = item.partition(sep)
        if found and 0 < len(head) <= 42 and rest:
            return head, rest
    return "", item


def outline_from_topic(topic: str, *, lang: str | None = "he", n_slides: int = 8, audience: str = "",
                       seed: int | None = None) -> list[Slide]:
    """A complete talk about `topic` with speaker notes.

    8 slides: title, agenda, background, key points, examples, plan, summary, Q&A. Longer decks
    add the challenge, a today-vs-goal comparison, a key-message quote, success metrics and
    section dividers; shorter ones keep the most important slides (3 to 14 slides).
    """
    lang = resolve_lang(lang, topic)
    t = _O[lang]
    subject = clean_presentation_topic(topic) or ("הנושא שלנו" if lang == "he" else "Our topic")
    title_text = subject[:1].upper() + subject[1:] if lang == "en" else subject
    domain = detect_domain(subject)
    business = _is_business(subject, domain)
    rng = make_rng("outline", subject, lang, audience, seed=seed)
    n = max(3, min(int(n_slides), len(_ORDER)))

    if audience:
        aud, aud_list = audience, audience
    elif business:
        auds = list(domain.audience[lang])
        aud, aud_list = auds[0], _join_list(auds[:3], lang)
    else:
        aud, aud_list = t["default_aud"], t["everyone"]
    to_aud = _he_prefix("ל", aud) if lang == "he" else aud
    topic_in_text = subject if lang == "he" else (subject if subject[:2].isupper() else subject[:1].lower() + subject[1:])

    def fill(s: str) -> str:
        return (s.replace("{aud_list}", aud_list).replace("{to_aud}", to_aud).replace("{aud}", aud)
                .replace("{topic}", topic_in_text).replace("{Topic}", title_text)
                .replace("{essence}", domain.t("essence", lang)).replace("{offer}", domain.t("offer", lang))
                .replace("{label}", domain.t("label", lang)))

    notes = {k: fill(v) for k, v in t["notes"].items()}
    plan = [fill(s) for s in t["plan_items"]]
    first_step = _head_text(plan[0])[1]

    if business:
        key_points = [f"{f.title[lang]}{' – ' if lang == 'he' else ': '}{f.text[lang]}" for f in domain.features]
        if domain.products:
            examples = [f"{p.name[lang]}{' – ' if lang == 'he' else ': '}{p.description[lang]}"
                        for p in domain.products[:3]]
        else:
            examples = [fill(s) for s in rng.sample(t["example_pool"], 3)]
        background = [fill(s) for s in t["background_business"]]
        sep = " – " if lang == "he" else ": "
        metrics = [f"{k}{sep}{how}" for k, how in zip(domain.kpis[lang][:3], t["kpi_tracking"], strict=False)]
        metrics.append(t["metrics_business_tail"])
    else:
        key_points = [fill(s) for s in [t["key_pool"][0], *rng.sample(t["key_pool"][1:], 3)]]
        examples = [fill(s) for s in rng.sample(t["example_pool"], 3)]
        background = [fill(s) for s in t["background_generic"]]
        metrics = [fill(s) for s in t["metrics_items"]]

    summary = [fill(s).replace("{next_step}", first_step) for s in t["summary_items"]]
    quote = rng.choice(t["quotes"])
    subtitle = t["for_audience"].replace("{audience}", audience) if audience else t["subtitle"]

    slides: dict[str, Slide] = {
        "title": Slide(title_text, [subtitle], notes["title"], "title"),
        "agenda": Slide(t["agenda"], [], notes["agenda"], "bullets", numbered=True),
        "section_1": Slide(t["section_1"], [title_text], notes["section"], "section"),
        "background": Slide(t["background"], background, notes["background"], "bullets"),
        "challenge": Slide(t["challenge"], [fill(s) for s in t["challenge_items"]], notes["challenge"], "bullets"),
        "key_points": Slide(t["key_points"], key_points, notes["key_points"], "bullets"),
        "compare": Slide(t["compare"], t["compare_today"] + t["compare_goal"], notes["compare"], "two_column",
                         columns=t["compare_cols"]),
        "examples": Slide(t["examples"], examples, notes["examples"], "bullets"),
        "quote": Slide(t["quote"], [quote, title_text], notes["quote"], "quote"),
        "section_2": Slide(t["section_2"], [title_text], notes["section"], "section"),
        "metrics": Slide(t["metrics"], metrics, notes["metrics"], "bullets"),
        "plan": Slide(t["plan"], plan, notes["plan"], "bullets", numbered=True),
        "summary": Slide(t["summary"], summary, notes["summary"], "bullets"),
        "qa": Slide(t["qa"], list(t["qa_lines"]), notes["qa"], "closing"),
    }
    keep = set(sorted(_ORDER, key=lambda k: _PRIORITY[k])[:n])
    chosen = [k for k in _ORDER if k in keep]
    agenda_items = [slides[k].title for k in chosen if k not in ("title", "agenda", "qa", "quote", "section_1",
                                                                  "section_2")]
    slides["agenda"].bullets = agenda_items[:8]
    return [slides[k] for k in chosen]


# --------------------------------------------------------------------------------------------
# HTML deck
# --------------------------------------------------------------------------------------------

_UI = {
    "he": {"next": "השקופית הבאה", "prev": "השקופית הקודמת", "notes": "הערות דובר", "fullscreen": "מסך מלא",
           "no_notes": "אין הערות לשקופית זו.",
           "hint": "חיצים / רווח – ניווט · N – הערות דובר · F – מסך מלא · Ctrl+P – הדפסה ל-PDF",
           "slide": "שקופית"},
    "en": {"next": "Next slide", "prev": "Previous slide", "notes": "Speaker notes", "fullscreen": "Full screen",
           "no_notes": "No notes for this slide.",
           "hint": "Arrows / Space: navigate · N: speaker notes · F: full screen · Ctrl+P: print to PDF",
           "slide": "Slide"},
}


def _bullet_html(item: str, idx: int, numbered: bool) -> str:
    head, text = _head_text(item)
    mark = f'<span class="mark num">{idx}</span>' if numbered else '<span class="mark"></span>'
    body = f"<b>{esc(head)}</b><span>{esc(text)}</span>" if head else f"<span>{esc(text)}</span>"
    return f'<li>{mark}<div class="txt">{body}</div></li>'


def _slide_html(slide: Slide, i: int, total: int, deck_title: str, section_no: int) -> str:
    notes = f'<aside class="notes">{esc(slide.notes)}</aside>' if slide.notes else ""
    foot = (f'<footer class="foot"><span>{esc(deck_title)}</span><span class="pg">{i} / {total}</span></footer>')
    layout = slide.layout
    if layout == "title":
        sub = "".join(f'<p class="sub">{esc(b)}</p>' for b in slide.bullets[:3])
        inner = (f'<div class="orbs" aria-hidden="true"><span></span><span></span><span></span></div>'
                 f'<div class="inner"><span class="kicker"></span><h1>{esc(slide.title)}</h1>{sub}</div>')
    elif layout == "closing":
        sub = "".join(f'<p class="sub">{esc(b)}</p>' for b in slide.bullets[:4])
        inner = (f'<div class="orbs" aria-hidden="true"><span></span><span></span><span></span></div>'
                 f'<div class="inner"><h1>{esc(slide.title)}</h1>{sub}</div>')
    elif layout == "section":
        sub = f'<p class="sub">{esc(slide.bullets[0])}</p>' if slide.bullets else ""
        inner = (f'<span class="big-no" aria-hidden="true">{section_no:02d}</span>'
                 f'<div class="inner"><h1>{esc(slide.title)}</h1>{sub}</div>{foot}')
    elif layout == "quote":
        quote = slide.bullets[0] if slide.bullets else slide.title
        by = slide.bullets[1] if len(slide.bullets) > 1 else ""
        cite = f'<figcaption>{esc(by)}</figcaption>' if by else ""
        inner = (f'<h2 class="eyebrow">{esc(slide.title)}</h2><figure><span class="qmark" aria-hidden="true">“</span>'
                 f'<blockquote>{esc(quote)}</blockquote>{cite}</figure>{foot}')
    elif layout == "two_column":
        half = (len(slide.bullets) + 1) // 2
        cols = [slide.bullets[:half], slide.bullets[half:]]
        heads = slide.columns or ("", "")
        cards = "".join(
            f'<div class="col c{k}">{f"<h3>{esc(heads[k])}</h3>" if heads[k] else ""}<ul class="bullets">'
            f'{"".join(_bullet_html(b, j + 1, slide.numbered) for j, b in enumerate(col))}</ul></div>'
            for k, col in enumerate(cols)
        )
        inner = f'<header class="head"><h2>{esc(slide.title)}</h2></header><div class="cols">{cards}</div>{foot}'
    else:
        many = " many" if len(slide.bullets) > 5 or sum(len(b) for b in slide.bullets) > 360 else ""
        items = "".join(_bullet_html(b, j + 1, slide.numbered) for j, b in enumerate(slide.bullets))
        inner = (f'<header class="head"><h2>{esc(slide.title)}</h2></header>'
                 f'<ul class="bullets{many}{" numbered" if slide.numbered else ""}">{items}</ul>{foot}')
    return (f'<section class="slide l-{layout}" data-layout="{layout}" id="s{i}" aria-roledescription="slide" '
            f'aria-label="{i} / {total}">{inner}{notes}</section>')


def _deck_css(palette: Palette) -> str:
    return f"""{palette.to_css()}
*{{box-sizing:border-box}}
html,body{{margin:0;height:100%;background:#0b0d12;overflow:hidden}}
body{{font-family:{FONT_STACK};color:var(--color-text);-webkit-font-smoothing:antialiased}}
.deck{{position:fixed;inset:0}}
.slide{{position:absolute;left:50%;top:50%;width:1280px;height:720px;margin:-360px 0 0 -640px;overflow:hidden;
transform:scale(var(--scale,1));background:var(--color-background);opacity:0;visibility:hidden;
transition:opacity .45s ease,visibility .45s;padding:72px 96px;display:flex;flex-direction:column}}
.slide.active{{opacity:1;visibility:visible}}
.slide .notes{{display:none}}
.head{{margin-bottom:36px}}
.head h2{{margin:0;font-size:50px;line-height:1.12;letter-spacing:-.015em;color:var(--color-text);position:relative;
padding-inline-start:28px}}
.head h2::before{{content:"";position:absolute;inset-inline-start:0;top:8px;bottom:8px;width:8px;border-radius:8px;
background:linear-gradient(var(--color-primary),var(--color-accent))}}
.bullets{{list-style:none;margin:0;padding:0;display:flex;flex-direction:column;gap:22px;font-size:30px;line-height:1.35}}
.bullets.many{{font-size:25px;gap:16px}}
.bullets li{{display:flex;gap:20px;align-items:flex-start}}
.mark{{flex:none;width:16px;height:16px;margin-top:.45em;border-radius:50%;background:var(--color-accent);
box-shadow:0 0 0 6px rgba(var(--color-accent-rgb),.18)}}
.mark.num{{width:46px;height:46px;margin-top:-2px;border-radius:14px;display:grid;place-items:center;font-size:22px;
font-weight:800;color:var(--color-on-primary);background:var(--color-primary);box-shadow:none}}
.txt b{{display:block;font-weight:800;color:var(--color-primary)}}
.txt span{{color:var(--color-text)}}
.txt b+span{{color:var(--color-muted);font-size:.86em}}
.foot{{margin-top:auto;display:flex;justify-content:space-between;font-size:16px;color:var(--color-muted);
border-top:1px solid var(--color-border);padding-top:16px}}
.l-title,.l-closing{{background:linear-gradient(135deg,var(--color-primary),var(--color-secondary));
color:var(--color-on-primary);justify-content:center}}
.orbs span{{position:absolute;border-radius:50%;filter:blur(2px)}}
.orbs span:nth-child(1){{width:520px;height:520px;inset-inline-end:-160px;top:-180px;
background:radial-gradient(circle at 30% 30%,rgba(var(--color-accent-rgb),.75),rgba(var(--color-accent-rgb),0) 70%)}}
.orbs span:nth-child(2){{width:300px;height:300px;inset-inline-end:240px;bottom:-150px;
background:radial-gradient(circle,rgba(255,255,255,.22),rgba(255,255,255,0) 70%)}}
.orbs span:nth-child(3){{width:180px;height:180px;inset-inline-start:-60px;top:80px;border:2px solid rgba(255,255,255,.25)}}
.l-title .inner,.l-closing .inner{{position:relative;max-width:900px}}
.l-title .kicker{{display:block;width:84px;height:8px;border-radius:8px;background:var(--color-accent);margin-bottom:34px}}
.l-title h1{{margin:0 0 22px;font-size:84px;line-height:1.04;letter-spacing:-.03em}}
.l-closing h1{{margin:0 0 22px;font-size:96px;line-height:1;letter-spacing:-.03em}}
.l-title .sub,.l-closing .sub{{margin:0 0 8px;font-size:30px;opacity:.9}}
.l-section{{background:var(--color-primary);color:var(--color-on-primary);justify-content:center}}
.l-section .big-no{{position:absolute;inset-inline-end:60px;bottom:-60px;font-size:420px;font-weight:900;line-height:1;
opacity:.12;letter-spacing:-.05em}}
.l-section h1{{margin:0 0 16px;font-size:68px;line-height:1.08;max-width:900px;position:relative}}
.l-section .sub{{margin:0;font-size:28px;opacity:.85;position:relative}}
.l-section .foot{{color:inherit;opacity:.7;border-top-color:rgba(255,255,255,.25);position:relative}}
.l-quote{{background:var(--color-surface);justify-content:center}}
.l-quote .eyebrow{{position:absolute;top:72px;margin:0;font-size:20px;letter-spacing:.12em;text-transform:uppercase;
color:var(--color-muted)}}
.l-quote figure{{margin:0;position:relative;max-width:1000px}}
.l-quote .qmark{{position:absolute;inset-inline-start:-30px;top:-120px;font-size:260px;line-height:1;font-family:Georgia,serif;
color:var(--color-accent);opacity:.55}}
.l-quote blockquote{{margin:0;font-size:58px;line-height:1.2;font-weight:800;letter-spacing:-.02em;position:relative}}
.l-quote figcaption{{margin-top:28px;font-size:24px;color:var(--color-muted)}}
.l-quote figcaption::before{{content:"— "}}
.l-quote .foot{{position:absolute;inset-inline:96px;bottom:56px}}
.cols{{display:grid;grid-template-columns:1fr 1fr;gap:32px;flex:1;min-height:0}}
.col{{background:var(--color-surface);border:1px solid var(--color-border);border-radius:28px;padding:34px 36px;
box-shadow:0 18px 40px rgba(0,0,0,.06)}}
.col h3{{margin:0 0 24px;font-size:22px;letter-spacing:.06em;text-transform:uppercase;color:var(--color-muted)}}
.col.c1{{background:linear-gradient(160deg,rgba(var(--color-primary-rgb),.1),rgba(var(--color-accent-rgb),.12));
border-color:rgba(var(--color-primary-rgb),.25)}}
.col.c1 h3{{color:var(--color-primary)}}
.col .bullets{{font-size:27px;gap:20px}}
.col.c0 .mark{{background:var(--color-muted);box-shadow:0 0 0 6px rgba(var(--color-muted-rgb),.15)}}
.progress{{position:fixed;left:0;right:0;bottom:0;height:5px;background:rgba(255,255,255,.08);z-index:5}}
.progress span{{display:block;height:100%;width:0;background:linear-gradient(90deg,var(--color-primary),var(--color-accent));
transition:width .35s ease}}
[dir=rtl] .progress span{{margin-inline-start:0;background:linear-gradient(270deg,var(--color-primary),var(--color-accent))}}
.controls{{position:fixed;bottom:16px;inset-inline-end:16px;display:flex;gap:8px;z-index:6;opacity:.25;
transition:opacity .2s}}
.controls:hover,.controls:focus-within{{opacity:1}}
.controls button{{font:600 14px/1 {FONT_STACK};min-width:42px;height:42px;padding:0 12px;border-radius:12px;
border:1px solid rgba(255,255,255,.2);background:rgba(20,22,28,.75);color:#fff;cursor:pointer}}
.controls button:focus-visible{{outline:3px solid var(--color-accent);outline-offset:2px}}
.counter{{position:fixed;bottom:22px;inset-inline-start:20px;font:600 14px/1 {FONT_STACK};color:rgba(255,255,255,.55);
z-index:6}}
.notes-panel{{position:fixed;inset-inline:16px;bottom:72px;max-width:900px;margin-inline:auto;padding:18px 22px;
border-radius:16px;background:rgba(15,17,23,.94);color:#f1f5f9;font-size:18px;line-height:1.55;z-index:7;
box-shadow:0 20px 50px rgba(0,0,0,.4)}}
.notes-panel h2{{margin:0 0 6px;font-size:13px;letter-spacing:.1em;text-transform:uppercase;color:#94a3b8}}
.notes-panel p{{margin:0}}
.hint{{position:fixed;top:16px;left:50%;transform:translateX(-50%);padding:10px 16px;border-radius:999px;
background:rgba(15,17,23,.85);color:#e2e8f0;font-size:14px;z-index:8;transition:opacity .6s;white-space:nowrap}}
.hint.gone{{opacity:0;pointer-events:none}}
@media (prefers-reduced-motion:reduce){{.slide{{transition:none}}}}
@page{{size:1280px 720px;margin:0}}
@media print{{
html,body{{height:auto;overflow:visible;background:none}}
.deck{{position:static}}
.slide{{position:relative;left:auto;top:auto;margin:0;transform:none!important;opacity:1!important;visibility:visible!important;
break-after:page;page-break-after:always;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.progress,.controls,.counter,.notes-panel,.hint{{display:none!important}}
}}
"""


_DECK_JS = r"""
(() => {
  const slides = Array.from(document.querySelectorAll('.slide'));
  const bar = document.querySelector('.progress span');
  const counter = document.querySelector('.counter');
  const panel = document.querySelector('.notes-panel');
  const panelText = panel.querySelector('p');
  const rtl = document.documentElement.dir === 'rtl';
  const total = slides.length;
  let current = 0;

  function fit() {
    const scale = Math.min(window.innerWidth / 1280, window.innerHeight / 720);
    document.documentElement.style.setProperty('--scale', String(scale));
  }

  function show(index, push) {
    current = Math.max(0, Math.min(total - 1, index));
    slides.forEach((s, i) => {
      const on = i === current;
      s.classList.toggle('active', on);
      s.setAttribute('aria-hidden', on ? 'false' : 'true');
    });
    bar.style.width = (total > 1 ? (current / (total - 1)) * 100 : 100) + '%';
    if (rtl) { bar.style.marginInlineStart = '0'; }
    counter.textContent = (current + 1) + ' / ' + total;
    const notes = slides[current].querySelector('.notes');
    panelText.textContent = notes ? notes.textContent : panel.dataset.empty;
    if (push !== false) { history.replaceState(null, '', '#' + (current + 1)); }
  }

  const next = () => show(current + 1);
  const prev = () => show(current - 1);
  const toggleNotes = () => { panel.hidden = !panel.hidden; };
  const fullscreen = () => {
    if (!document.fullscreenElement) { document.documentElement.requestFullscreen?.(); }
    else { document.exitFullscreen?.(); }
  };

  document.addEventListener('keydown', (e) => {
    if (e.ctrlKey || e.metaKey || e.altKey) { return; }
    const k = e.key;
    if (k === 'ArrowRight' || k === 'ArrowDown' || k === 'PageDown' || k === ' ' || k === 'Enter') { e.preventDefault(); next(); }
    else if (k === 'ArrowLeft' || k === 'ArrowUp' || k === 'PageUp' || k === 'Backspace') { e.preventDefault(); prev(); }
    else if (k === 'Home') { show(0); }
    else if (k === 'End') { show(total - 1); }
    else if (k === 'n' || k === 'N' || k === 'מ') { toggleNotes(); }
    else if (k === 'f' || k === 'F' || k === 'כ') { fullscreen(); }
  });

  document.querySelector('[data-act=next]').addEventListener('click', next);
  document.querySelector('[data-act=prev]').addEventListener('click', prev);
  document.querySelector('[data-act=notes]').addEventListener('click', toggleNotes);
  document.querySelector('[data-act=full]').addEventListener('click', fullscreen);

  let touchX = null;
  document.addEventListener('touchstart', (e) => { touchX = e.touches[0].clientX; }, { passive: true });
  document.addEventListener('touchend', (e) => {
    if (touchX === null) { return; }
    const dx = e.changedTouches[0].clientX - touchX;
    touchX = null;
    if (Math.abs(dx) < 40) { return; }
    const forward = rtl ? dx > 0 : dx < 0;
    if (forward) { next(); } else { prev(); }
  });

  window.addEventListener('resize', fit);
  window.addEventListener('hashchange', () => {
    const n = parseInt(location.hash.slice(1), 10);
    if (!Number.isNaN(n)) { show(n - 1, false); }
  });

  fit();
  const start = parseInt(location.hash.slice(1), 10);
  show(Number.isNaN(start) ? 0 : start - 1, false);
  const hint = document.querySelector('.hint');
  setTimeout(() => hint.classList.add('gone'), 4500);
})();
"""


def build_presentation_html(title: str, slides: Sequence[Slide], *, palette: Palette | None = None,
                            lang: str | None = "he") -> str:
    """A single self-contained HTML slide deck (no external files).

    Keys: arrows / space / PageUp / PageDown navigate, Home / End jump, N toggles speaker
    notes, F toggles full screen; swipe on touch screens. Ctrl+P prints one slide per page.
    """
    lang = resolve_lang(lang, title)
    if not slides:
        raise ValueError("a presentation needs at least one slide")
    palette = palette or generate_palette(title)
    ui = _UI[lang]
    total = len(slides)
    section_no = 0
    parts = []
    for i, slide in enumerate(slides, 1):
        if slide.layout == "section":
            section_no += 1
        parts.append(_slide_html(slide, i, total, title, max(1, section_no)))
    arrow_next, arrow_prev = ("←", "→") if lang == "he" else ("→", "←")
    return f"""<!DOCTYPE html>
<html lang="{lang}" dir="{text_dir(lang)}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="generator" content="JARVIS presentation">
<title>{esc(title)}</title>
<style>{_deck_css(palette)}</style>
</head>
<body>
<main class="deck" aria-live="polite">
{"".join(parts)}
</main>
<div class="progress" aria-hidden="true"><span></span></div>
<div class="counter" aria-hidden="true">1 / {total}</div>
<nav class="controls" aria-label="{esc(ui["slide"])}">
<button type="button" data-act="prev" aria-label="{esc(ui["prev"])}">{arrow_prev}</button>
<button type="button" data-act="next" aria-label="{esc(ui["next"])}">{arrow_next}</button>
<button type="button" data-act="notes" aria-label="{esc(ui["notes"])}">N</button>
<button type="button" data-act="full" aria-label="{esc(ui["fullscreen"])}">⛶</button>
</nav>
<section class="notes-panel" hidden data-empty="{esc(ui["no_notes"])}"><h2>{esc(ui["notes"])}</h2><p></p></section>
<div class="hint" role="note">{esc(ui["hint"])}</div>
<script>{_DECK_JS}</script>
</body>
</html>
"""


# --------------------------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------------------------

_PPT_FONT = "Segoe UI"


def _rgb(hex_color: str):  # -> pptx.dml.color.RGBColor
    from pptx.dml.color import RGBColor

    return RGBColor(*hex_to_rgb(hex_color))


class _PptxWriter:
    """Draws slides with python-pptx on the built-in layouts (so titles and bodies stay real,
    editable placeholders), adding themed background shapes behind them."""

    W = 12192000  # 13.333 in
    H = 6858000  # 7.5 in

    def __init__(self, title: str, palette: Palette, lang: str) -> None:
        from pptx import Presentation

        self.title = title
        self.p = palette
        self.lang = lang
        self.rtl = lang == "he"
        self.prs = Presentation()
        self.prs.slide_width = self.W
        self.prs.slide_height = self.H
        props = self.prs.core_properties
        props.title = title
        props.author = "JARVIS"
        props.language = "he-IL" if self.rtl else "en-US"

    # ----- low-level helpers --------------------------------------------------------------
    def _stack_back(self, slide, shapes: Sequence) -> None:
        """Move `shapes` (listed bottom to top) beneath everything else on the slide."""
        tree = slide.shapes._spTree
        for k, shp in enumerate(shapes):
            tree.remove(shp._element)
            tree.insert(2 + k, shp._element)

    def _rect(self, slide, x, y, w, h, color: str, *, shape: str = "rect", alpha: float | None = None,
              line: str | None = None, gradient: tuple[str, str] | None = None, angle: float = 0):
        from pptx.enum.shapes import MSO_SHAPE

        kind = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL}[shape]
        shp = slide.shapes.add_shape(kind, int(x), int(y), int(w), int(h))
        if shape == "round":
            shp.adjustments[0] = 0.08
        if gradient:
            shp.fill.gradient()
            shp.fill.gradient_angle = angle
            stops = shp.fill.gradient_stops
            stops[0].color.rgb = _rgb(gradient[0])
            stops[0].position = 0.0
            stops[1].color.rgb = _rgb(gradient[1])
            stops[1].position = 1.0
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = _rgb(color)
            if alpha is not None:
                self._alpha(shp, alpha)
        if line:
            shp.line.color.rgb = _rgb(line)
            shp.line.width = 12700
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    @staticmethod
    def _alpha(shape, alpha: float) -> None:
        from pptx.oxml.ns import qn

        srgb = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        el = srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
        srgb.append(el)

    def _style_run(self, run, *, size: float, color: str, bold: bool = False, italic: bool = False) -> None:
        from pptx.oxml.ns import qn
        from pptx.util import Pt

        f = run.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.name = _PPT_FONT
        f.color.rgb = _rgb(color)
        rpr = run._r.get_or_add_rPr()
        rpr.set("lang", "he-IL" if has_hebrew(run.text) else "en-US")
        if has_hebrew(run.text):
            rpr.set("altLang", "en-US")
        for tag in ("a:cs",):
            if rpr.find(qn(tag)) is None:
                cs = rpr.makeelement(qn(tag), {"typeface": _PPT_FONT})
                latin = rpr.find(qn("a:latin"))
                if latin is not None:
                    latin.addnext(cs)
                else:
                    rpr.append(cs)

    def _para_dir(self, para, *, align: str | None = None) -> None:
        from pptx.enum.text import PP_ALIGN

        ppr = para._p.get_or_add_pPr()
        if self.rtl:
            ppr.set("rtl", "1")
        target = align or ("right" if self.rtl else "left")
        para.alignment = {"right": PP_ALIGN.RIGHT, "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[target]

    def _fill_text(self, tf, lines: Sequence[str], *, size: float, color: str, bold: bool = False,
                   align: str | None = None, space_after: float = 6, bullets: bool = False,
                   bullet_color: str | None = None, numbered: bool = False, head_color: str | None = None,
                   anchor: str | None = None) -> None:
        from pptx.enum.text import MSO_ANCHOR
        from pptx.oxml.ns import qn
        from pptx.util import Pt

        tf.word_wrap = True
        if anchor:
            tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        body_pr = tf._txBody.find(qn("a:bodyPr"))
        if body_pr is not None and body_pr.find(qn("a:normAutofit")) is None:
            for child in list(body_pr):
                if child.tag in (qn("a:spAutoFit"), qn("a:noAutofit")):
                    body_pr.remove(child)
            body_pr.append(body_pr.makeelement(qn("a:normAutofit"), {}))
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.space_after = Pt(space_after)
            self._para_dir(para, align=align)
            ppr = para._p.get_or_add_pPr()
            if bullets:
                indent = 457200 if numbered else 342900
                ppr.set("marL" if not self.rtl else "marR", str(indent))
                ppr.set("indent", str(-indent))
                if bullet_color:
                    clr = ppr.makeelement(qn("a:buClr"), {})
                    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": bullet_color.lstrip("#").upper()}))
                    ppr.append(clr)
                ppr.append(ppr.makeelement(qn("a:buSzPct"), {"val": "100000"}))
                if numbered:
                    ppr.append(ppr.makeelement(qn("a:buFont"), {"typeface": "+mj-lt"}))
                    scheme = "arabicPeriod" if not self.rtl else "arabicParenR"
                    ppr.append(ppr.makeelement(qn("a:buAutoNum"), {"type": scheme}))
                else:
                    ppr.append(ppr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
                    ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "●"}))
            else:
                ppr.append(ppr.makeelement(qn("a:buNone"), {}))
            head, text = _head_text(line) if head_color else ("", line)
            if head:
                r1 = para.add_run()
                r1.text = head + (" – " if self.rtl else ": ")
                self._style_run(r1, size=size, color=head_color, bold=True)
                r2 = para.add_run()
                r2.text = text
                self._style_run(r2, size=size, color=color, bold=bold)
            else:
                run = para.add_run()
                run.text = line
                self._style_run(run, size=size, color=color, bold=bold)

    def _place(self, shape, x, y, w, h) -> None:
        shape.left, shape.top, shape.width, shape.height = int(x), int(y), int(w), int(h)

    def _ph(self, slide, idx: int):
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == idx:
                return ph
        return None

    def _drop_unused(self, slide, used: set[int]) -> None:
        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx not in used:
                ph._element.getparent().remove(ph._element)

    def _textbox(self, slide, x, y, w, h, lines: Sequence[str], **kw):
        box = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
        self._fill_text(box.text_frame, lines, **kw)
        return box

    def _footer(self, slide, i: int, total: int, *, color: str) -> None:
        m = 640000
        y = self.H - 520000
        num_align = "left" if self.rtl else "right"
        title_align = "right" if self.rtl else "left"
        self._textbox(slide, m, y, self.W - 2 * m, 300000, [self.title], size=11, color=color, align=title_align,
                      space_after=0)
        self._textbox(slide, m, y, self.W - 2 * m, 300000, [f"{i} / {total}"], size=11, color=color, bold=True,
                      align=num_align, space_after=0)

    def _notes(self, slide, notes: str) -> None:
        if not notes:
            return
        tf = slide.notes_slide.notes_text_frame
        tf.text = notes
        for para in tf.paragraphs:
            self._para_dir(para)

    # ----- slide kinds ----------------------------------------------------------------------
    def _hero(self, slide, s: Slide, *, closing: bool) -> None:
        p = self.p
        bg = self._rect(slide, 0, 0, self.W, self.H, p.primary, gradient=(p.primary, p.secondary), angle=315)
        end_x = 0 if self.rtl else self.W  # decorations sit on the side where reading ends
        big = self._rect(slide, end_x - 2200000, -1500000, 4400000, 4400000, p.accent, shape="oval", alpha=0.35)
        small_x = (self.W * 0.62) if self.rtl else (self.W * 0.38 - 2600000)
        small = self._rect(slide, small_x, self.H - 1300000, 2600000, 2600000, "#ffffff", shape="oval", alpha=0.10)
        self._stack_back(slide, [bg, big, small])
        ink = p.on_primary
        margin = 900000
        title_ph, sub_ph = self._ph(slide, 0), self._ph(slide, 1)
        self._drop_unused(slide, {0, 1})
        if not closing:
            self._rect(slide, (self.W - margin - 760000) if self.rtl else margin, 2000000, 760000, 76000, p.accent,
                       shape="round")
        self._place(title_ph, margin, 2150000 if not closing else 1900000, self.W - 2 * margin, 1800000)
        title_ph.text_frame.clear()
        self._fill_text(title_ph.text_frame, [s.title], size=54 if not closing else 60, color=ink, bold=True,
                        anchor="bottom" if not closing else "middle", space_after=0)
        self._place(sub_ph, margin, 4050000, self.W - 2 * margin, 1300000)
        sub_ph.text_frame.clear()
        self._fill_text(sub_ph.text_frame, s.bullets[:4] or [" "], size=24, color=ink, anchor="top")

    def _section(self, slide, s: Slide, i: int, total: int, section_no: int) -> None:
        p = self.p
        bg = self._rect(slide, 0, 0, self.W, self.H, p.primary)
        big = self._textbox(slide, 0 if self.rtl else self.W - 5600000, self.H - 4300000, 5600000, 4300000,
                            [f"{section_no:02d}"], size=250, color=mix(p.primary, p.on_primary, 0.16), bold=True,
                            align="left" if self.rtl else "right", space_after=0, anchor="bottom")
        self._stack_back(slide, [bg, big])
        title_ph, body_ph = self._ph(slide, 0), self._ph(slide, 1)
        self._drop_unused(slide, {0, 1})
        margin = 900000
        width = self.W - 2 * margin - 1500000
        x = margin + 1500000 if self.rtl else margin
        self._place(title_ph, x, 2300000, width, 1500000)
        title_ph.text_frame.clear()
        self._fill_text(title_ph.text_frame, [s.title], size=44, color=p.on_primary, bold=True, anchor="bottom",
                        space_after=0)
        self._place(body_ph, x, 3900000, width, 900000)
        body_ph.text_frame.clear()
        self._fill_text(body_ph.text_frame, s.bullets[:1] or [" "], size=22, color=p.on_primary, anchor="top")
        self._footer(slide, i, total, color=mix(p.primary, p.on_primary, 0.7))

    def _content_frame(self, slide, s: Slide) -> list:
        """Background, top stripe, title placeholder and its accent bar; returns shapes to stack."""
        p = self.p
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(p.background)
        stripe = self._rect(slide, 0, 0, self.W, 90000, p.primary, gradient=(p.primary, p.accent), angle=0)
        margin = 700000
        title_ph = self._ph(slide, 0)
        self._place(title_ph, margin if self.rtl else margin + 200000, 420000, self.W - 2 * margin - 200000, 1000000)
        title_ph.text_frame.clear()
        self._fill_text(title_ph.text_frame, [s.title], size=34, color=p.text, bold=True, anchor="middle",
                        space_after=0)
        bar_x = (self.W - margin - 76000) if self.rtl else margin
        self._rect(slide, bar_x, 600000, 76000, 640000, p.primary, shape="round", gradient=(p.primary, p.accent),
                   angle=90)
        return [stripe]

    def _bullets(self, slide, s: Slide, i: int, total: int) -> None:
        p = self.p
        self._stack_back(slide, self._content_frame(slide, s))
        body = self._ph(slide, 1)
        self._drop_unused(slide, {0, 1})
        margin = 700000
        self._place(body, margin, 1600000, self.W - 2 * margin, 4500000)
        body.text_frame.clear()
        n_chars = sum(len(b) for b in s.bullets)
        size = 24 if len(s.bullets) <= 4 and n_chars < 260 else (20 if len(s.bullets) <= 6 and n_chars < 420 else 17)
        self._fill_text(body.text_frame, s.bullets or [" "], size=size, color=p.text, bullets=True,
                        bullet_color=p.accent if not s.numbered else p.primary, numbered=s.numbered,
                        head_color=p.primary, space_after=14, anchor="top")
        self._footer(slide, i, total, color=p.muted)

    def _two_column(self, slide, s: Slide, i: int, total: int) -> None:
        p = self.p
        back = self._content_frame(slide, s)
        margin, gap, top, h = 700000, 380000, 1650000, 4300000
        col_w = (self.W - 2 * margin - gap) / 2
        half = (len(s.bullets) + 1) // 2
        cols = [s.bullets[:half], s.bullets[half:]]
        heads = s.columns or ("", "")
        xs = [margin, margin + col_w + gap]
        if self.rtl:  # the first column starts on the right
            xs.reverse()
        tints = [mix(p.surface, p.muted, 0.06), mix(p.background, p.primary, 0.10)]
        lines = [p.border, mix(p.border, p.primary, 0.3)]
        cards = [self._rect(slide, xs[k], top, col_w, h, tints[k], shape="round", line=lines[k]) for k in range(2)]
        self._stack_back(slide, back + cards)
        pad = 300000
        for k, (head_idx, body_idx) in enumerate(((1, 2), (3, 4))):
            head_ph, body_ph = self._ph(slide, head_idx), self._ph(slide, body_idx)
            self._place(head_ph, xs[k] + pad, top + 220000, col_w - 2 * pad, 560000)
            head_ph.text_frame.clear()
            self._fill_text(head_ph.text_frame, [heads[k] or " "], size=18, color=p.primary if k else p.muted,
                            bold=True, space_after=0, anchor="middle")
            self._place(body_ph, xs[k] + pad, top + 850000, col_w - 2 * pad, h - 1050000)
            body_ph.text_frame.clear()
            self._fill_text(body_ph.text_frame, cols[k] or [" "], size=20, color=p.text, bullets=True,
                            bullet_color=p.primary if k else p.muted, space_after=12, anchor="top")
        self._drop_unused(slide, {0, 1, 2, 3, 4})
        self._footer(slide, i, total, color=p.muted)

    def _quote(self, slide, s: Slide, i: int, total: int) -> None:
        p = self.p
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(p.surface)
        title_ph = self._ph(slide, 0)
        self._drop_unused(slide, {0})
        margin = 1100000
        self._place(title_ph, margin, 500000, self.W - 2 * margin, 600000)
        title_ph.text_frame.clear()
        self._fill_text(title_ph.text_frame, [s.title], size=16, color=p.muted, bold=True, space_after=0)
        mark_x = (self.W - margin - 1300000) if self.rtl else margin - 300000
        mark = self._textbox(slide, mark_x, 1000000, 1600000, 1800000, ["“"], size=180,
                             color=mix(p.surface, p.accent, 0.7), bold=True, align="right" if self.rtl else "left",
                             space_after=0)
        self._stack_back(slide, [mark])
        quote = s.bullets[0] if s.bullets else s.title
        self._textbox(slide, margin, 2050000, self.W - 2 * margin, 2300000, [quote], size=38, color=p.text, bold=True,
                      anchor="middle", space_after=0)
        if len(s.bullets) > 1:
            self._textbox(slide, margin, 4450000, self.W - 2 * margin, 600000, [f"— {s.bullets[1]}"], size=20,
                          color=p.muted, space_after=0)
        self._footer(slide, i, total, color=p.muted)

    # ----- public ----------------------------------------------------------------------------
    def add(self, s: Slide, i: int, total: int, section_no: int) -> None:
        layouts = self.prs.slide_layouts
        index = {"title": 0, "closing": 0, "section": 2, "bullets": 1, "two_column": 4, "quote": 5}[s.layout]
        slide = self.prs.slides.add_slide(layouts[index])
        if s.layout in ("title", "closing"):
            self._hero(slide, s, closing=s.layout == "closing")
        elif s.layout == "section":
            self._section(slide, s, i, total, section_no)
        elif s.layout == "two_column":
            self._two_column(slide, s, i, total)
        elif s.layout == "quote":
            self._quote(slide, s, i, total)
        else:
            self._bullets(slide, s, i, total)
        self._notes(slide, s.notes)

    def save(self, target) -> None:
        self.prs.save(target)


def _write_pptx(title: str, slides: Sequence[Slide], target, palette: Palette, lang: str) -> None:
    writer = _PptxWriter(title, palette, lang)
    total = len(slides)
    section_no = 0
    for i, s in enumerate(slides, 1):
        if s.layout == "section":
            section_no += 1
        writer.add(s, i, total, max(1, section_no))
    writer.save(target)


def build_presentation_pptx(title: str, slides: Sequence[Slide], path: str | Path, *, palette: Palette | None = None,
                            lang: str | None = "he") -> Path:
    """Write an editable 16:9 PowerPoint deck (titles and bullets are real placeholders, notes
    included). Hebrew paragraphs are right-aligned with rtl="1". Returns the path."""
    lang = resolve_lang(lang, title)
    if not slides:
        raise ValueError("a presentation needs at least one slide")
    palette = palette or generate_palette(title)
    path = Path(path)
    if path.suffix.lower() != ".pptx":
        path = path.with_suffix(".pptx")
    path.parent.mkdir(parents=True, exist_ok=True)
    _write_pptx(title, slides, path, palette, lang)
    return path


def generate_presentation(topic: str, *, lang: str | None = None, n_slides: int = 8, audience: str = "",
                          palette: Palette | None = None, slides: Sequence[Slide] | None = None,
                          title: str | None = None) -> list[Asset]:
    """Outline + HTML deck + PPTX deck as assets (nothing is written to disk)."""
    lang = resolve_lang(lang, topic)
    slides = list(slides) if slides else outline_from_topic(topic, lang=lang, n_slides=n_slides, audience=audience)
    deck_title = title or (slides[0].title if slides and slides[0].layout == "title" else clean_presentation_topic(topic))
    palette = palette or generate_palette(topic)
    buf = io.BytesIO()
    _write_pptx(deck_title, slides, buf, palette, lang)
    outline = json.dumps({"title": deck_title, "lang": lang, "slides": [s.to_dict() for s in slides]},
                         ensure_ascii=False, indent=2)
    label = "מצגת" if lang == "he" else "presentation"
    return [
        Asset(name=f"{label}-{deck_title}", kind="html",
              content=build_presentation_html(deck_title, slides, palette=palette, lang=lang),
              description="HTML deck: arrows to navigate, N for notes, Ctrl+P for PDF" if lang == "en"
              else "מצגת HTML: חיצים לניווט, N להערות, Ctrl+P ל-PDF"),
        Asset(name=f"{label}-{deck_title}", kind="pptx", content=buf.getvalue(),
              description="PowerPoint 16:9 with speaker notes" if lang == "en" else "PowerPoint 16:9 עם הערות דובר"),
        Asset(name=f"outline-{deck_title}", kind="json", content=outline,
              description="Slide outline (JSON)" if lang == "en" else "מבנה המצגת (JSON)"),
    ]
